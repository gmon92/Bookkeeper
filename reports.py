"""Deterministic USD reports from reviewed SQLite transactions."""
import io
import json
import base64
import hashlib
import binascii
import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal
from credits import LABELS as CREDIT_LABELS

PROFILE_SHEET = '_Bookkeeper_Profile'
PROFILE_FORMAT = 'local-bookkeeper-workbook-profile'
MAX_PROFILE_BYTES = 64 * 1024 * 1024
PROFILE_CHUNK_SIZE = 30000
TRANSACTION_HEADERS = ['Date','Description','Business','Source account','Category','Amount USD','Direction',
    'Net expense USD','Review status','Classification','Included in totals','Outside totals reason',
    'Merchant','Posted date','Kind','Excluded','Credit decision','Statement category','Source document',
    'Source sheet','Source row','Notes','Transaction ID','Related transaction ID','Duplicate of']


def transaction_values(row):
    from source_archive import category as statement_category
    row=dict(row)
    raw=json.loads(row.get('source_raw_json') or '{}')
    account=row.get('nickname') or ''
    if row.get('last_four'):account+=' ···· '+row['last_four']
    statuses={'approved':'Reviewed','needs_review':'Awaiting review','suggested':'Categorized / awaiting review'}
    return [row['date'],row['description'],row.get('business_name') or '',account,row['category_name'],
        Decimal(row['amount_cents'])/100,row['direction'],Decimal(expense_cents(row))/100 if included(row) else None,
        statuses.get(row['status'],row['status']),row['classification'],bool(included(row)),outside_totals_reason(row),
        row.get('merchant',''),row.get('posted_date'),row['transaction_kind'],bool(row['excluded']),
        CREDIT_LABELS.get(row['credit_treatment'],'') if row['direction']=='credit' or row['credit_treatment'] else '',
        statement_category(raw),row.get('filename') or '',raw.get('sheet',''),row.get('source_row'),row.get('notes',''),
        row['id'],row.get('related_transaction_id'),row['duplicate_of']]


def csv_transactions(rows):
    output=io.StringIO(newline='')
    writer=csv.writer(output)
    writer.writerow(TRANSACTION_HEADERS)
    for row in rows:
        writer.writerow([safe_text(v) if isinstance(v,str) else format(v,'.2f') if isinstance(v,Decimal) else v
            for v in transaction_values(row)])
    return output.getvalue().encode('utf-8-sig')


def add_profile(workbook, database, profile):
    if len(database)>MAX_PROFILE_BYTES:
        raise ValueError('This profile is too large for an embedded Excel copy. Export SQLite for complete books, or use the accountant summary without profile data.')
    encoded = base64.b64encode(database).decode('ascii')
    chunks = [encoded[i:i+PROFILE_CHUNK_SIZE] for i in range(0,len(encoded),PROFILE_CHUNK_SIZE)]
    manifest = {'format':PROFILE_FORMAT,'version':1,'bytes':len(database),'sha256':hashlib.sha256(database).hexdigest(),'chunks':len(chunks)}
    payload = workbook.create_sheet(PROFILE_SHEET)
    payload.append([PROFILE_FORMAT]); payload.append([json.dumps(manifest,sort_keys=True)])
    for index,chunk in enumerate(chunks):
        payload.append([index,chunk]); payload.cell(payload.max_row,2).data_type='s'
    payload.sheet_state='veryHidden'
    info = workbook.create_sheet('Profile Restore')
    for row in [
        ['Reusable Local Bookkeeper profile'],
        ['Scope','Full active profile, including ALL businesses and years; independent of the report filters.'],
        ['Profile name',safe_text(profile.get('name','My books'))],
        ['Transactions',profile.get('transactionCount',0)],['Saved rules','Import separately from SQLite or a rules JSON file.'],
        ['What is retained','Transactions, categorizations, credit decisions, account scopes, categories, businesses, import layouts, preferences and original row archives.'],
        ['Resume','In Local Bookkeeper, open the profile chooser, choose Restore export / backup, select this Excel file, then choose a new empty destination.'],
        ['Original statement files','Use Download profile backup to also transfer the original statement files. Original rows remain inside this profile.'],
        ['Report edits','Editing report cells does not change the embedded profile snapshot. Export a fresh copy after making decisions in the app.'],
        ['Integrity','Restore validates the embedded database checksum and SQLite integrity before opening.'],
        ['Rules','Excel and ZIP imports do not restore automation rules. Use an exported SQLite database or Import JSON in Rules & automation.'],
    ]: info.append(row)
    info.column_dimensions['A'].width=26; info.column_dimensions['B'].width=110
    from openpyxl.styles import Alignment, Font
    for row in info:
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
        info.row_dimensions[row[0].row].height=42
    info['A1'].font=Font(bold=True,size=15)


def read_profile(workbook):
    if PROFILE_SHEET not in workbook.sheetnames:
        raise ValueError('This Excel file does not contain a reusable profile. Accountant summaries cannot restore books. Open an exported SQLite database or restore a full profile backup.')
    data = workbook[PROFILE_SHEET].iter_rows(values_only=True)
    if next(data,())[0:1] != (PROFILE_FORMAT,): raise ValueError('The workbook profile format is invalid.')
    try:
        manifest = json.loads(next(data)[0])
        if not isinstance(manifest,dict) or manifest.get('format')!=PROFILE_FORMAT or manifest.get('version')!=1: raise ValueError()
        size,count=manifest['bytes'],manifest['chunks']
        if type(size) is not int or not 0<size<=MAX_PROFILE_BYTES or type(count) is not int or not 0<count<=((MAX_PROFILE_BYTES*4//3+4)//PROFILE_CHUNK_SIZE+1): raise ValueError()
        parts=[]
        for index in range(count):
            row=next(data)
            if len(row)<2 or row[0]!=index or not isinstance(row[1],str) or len(row[1])>PROFILE_CHUNK_SIZE: raise ValueError()
            parts.append(row[1])
        if next(data,None) is not None: raise ValueError()
        content=base64.b64decode(''.join(parts),validate=True)
        if len(content)!=size or hashlib.sha256(content).hexdigest()!=manifest['sha256'] or not content.startswith(b'SQLite format 3\x00'): raise ValueError()
        return content
    except (ValueError,KeyError,TypeError,StopIteration,binascii.Error) as exc:
        raise ValueError('The reusable profile data is incomplete or its checksum does not match. Export another copy or use the SQLite database.') from exc


def safe_text(value):
    text = str(value or '')
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t','\r','\n')) else text


def included(row):
    return bool(row['include_in_totals'] and not row['excluded'] and row['status'] == 'approved' and row['classification'] == 'business' and row['category_name'].strip().casefold() != 'personal / non-business' and not row['duplicate_of'] and ((row['direction'] == 'debit' and row['transaction_kind'] in {'expense','fee'}) or (row['direction'] == 'credit' and row['transaction_kind'] == 'refund')))


def expense_cents(row):
    return -row['amount_cents'] if row['direction'] == 'credit' else row['amount_cents']


def outside_totals_reason(row):
    if included(row):return ''
    if row['duplicate_of']:return 'Duplicates'
    if row['classification']=='transfer' or row['transaction_kind']=='transfer' or row['credit_treatment']=='payment':return 'Transfers and card payments'
    if row['credit_treatment']=='statement_credit':return 'Statement credits and benefits'
    if row['classification']=='personal' or row['category_name'].strip().casefold()=='personal / non-business':return 'Personal activity'
    if row['excluded']:return 'Explicitly excluded activity'
    if row['status']!='approved':return 'Awaiting review'
    return 'Reviewed activity outside expense totals'


def accountant_summary(rows, business_label, year):
    """A filtered expense snapshot from saved inclusion decisions."""
    rows=[dict(r) for r in rows]
    categories,monthly=defaultdict(lambda:[0,0,0]),defaultdict(lambda:[0,0,0])
    outside=defaultdict(lambda:[0,0,0])
    charges=refunds=count=0
    for row in rows:
        if included(row):
            debit=row['amount_cents'] if row['direction']=='debit' else 0
            credit=row['amount_cents'] if row['direction']=='credit' else 0
            charges+=debit;refunds+=credit;count+=1
            for group,key in ((categories,(row['date'][:4],row.get('business_name') or business_label,row['category_name'])),
                              (monthly,(row['date'][:7],row.get('business_name') or business_label))):
                group[key][0]+=debit;group[key][1]+=credit;group[key][2]+=1
            continue
        reason=outside_totals_reason(row)
        outside[reason][0]+=1
        if row['direction']=='debit':outside[reason][1]+=row['amount_cents']
        elif row['direction']=='credit':outside[reason][2]+=row['amount_cents']
    dates=[r['date'] for r in rows if r['date']]
    return {'business':business_label,'period':year or 'All recorded years',
        'firstDate':min(dates) if dates else None,'lastDate':max(dates) if dates else None,
        'includedCount':count,'chargesCents':charges,'refundsCents':refunds,'netCents':charges-refunds,
        'pendingCount':sum(r['status']!='approved' for r in rows),
        'sourceAccountCount':len({r['account_id'] for r in rows}),
        'sourceDocumentCount':len({r['document_id'] for r in rows if r['document_id']}),
        'categories':[[*key,values[0],values[1],values[0]-values[1],values[2]] for key,values in sorted(categories.items())],
        'monthly':[[*key,values[0],values[1],values[0]-values[1],values[2]] for key,values in sorted(monthly.items())],
        'outsideTotals':[[key,*values] for key,values in sorted(outside.items())]}


def excel_summary(rows, business_label, year):
    # This recurring application export uses the app's existing Excel runtime.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows=[dict(row) for row in rows]
    data=accountant_summary(rows,business_label,year)
    workbook=Workbook();summary=workbook.active;summary.title='Accountant Summary'
    monthly=workbook.create_sheet('Monthly Summary')
    money_format='#,##0.00;[Red](#,##0.00);"—"'
    def money(cents):return Decimal(cents)/100
    def write(tab,row,values):
        for column,value in enumerate(values,1):
            tab.cell(row,column,safe_text(value) if isinstance(value,str) else value)
    def band(tab,row,values):
        write(tab,row,values)
        for cell in tab[row][:len(values)]:
            cell.fill=PatternFill('solid',fgColor='384D38');cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        tab.row_dimensions[row].height=30
    for tab,title in ((summary,'Expense summary'),(monthly,'Monthly expenses')):
        last_column='G' if tab==summary else 'F'
        tab.merge_cells('A1:'+last_column+'1');tab.merge_cells('A2:'+last_column+'2')
        write(tab,1,[title]);write(tab,2,[data['business']]);write(tab,3,['Reporting period',data['period']])
        tab.sheet_view.showGridLines=False
        tab.sheet_properties.pageSetUpPr.fitToPage=True
        tab.page_setup.orientation='landscape';tab.page_setup.paperSize=tab.PAPERSIZE_A4
        tab.page_setup.fitToWidth=1;tab.page_setup.fitToHeight=0
        tab.sheet_properties.outlinePr.summaryRight=False
        for column,width in enumerate((14,36,31,19,19,20,16),1):tab.column_dimensions[get_column_letter(column)].width=width
    write(summary,4,['Recorded activity',data['firstDate'] or 'No activity',data['lastDate'] or ''])
    write(summary,6,['Included charges',money(data['chargesCents'])]);summary['B6'].number_format=money_format
    write(summary,7,['Included refunds',money(data['refundsCents'])]);summary['B7'].number_format=money_format
    write(summary,8,['Net expenses',money(data['netCents'])]);summary['B8'].number_format=money_format
    write(summary,9,['Included rows',data['includedCount']]);write(summary,10,['Awaiting review',data['pendingCount']])
    write(summary,11,['Source accounts',data['sourceAccountCount'],None,'Source statements',data['sourceDocumentCount']])
    band(summary,13,['Year','Business','Category','Charges USD','Refunds USD','Net expense USD','Transactions'])
    for row,record in enumerate(data['categories'],14):write(summary,row,record[:3]+[money(v) for v in record[3:6]]+record[6:])
    total_row=14+len(data['categories'])
    write(summary,total_row,['TOTAL','','',money(data['chargesCents']),money(data['refundsCents']),money(data['netCents']),data['includedCount']])
    outside_row=total_row+3
    band(summary,outside_row,['Outside expense totals','Transactions','Money out USD','Money in USD'])
    for row,record in enumerate(data['outsideTotals'],outside_row+1):write(summary,row,record[:2]+[money(v) for v in record[2:]])
    note_row=outside_row+len(data['outsideTotals'])+3
    summary.merge_cells(start_row=note_row,start_column=1,end_row=note_row,end_column=7)
    write(summary,note_row,['Saved expense decisions only. Pending, personal, transfer and excluded activity stays outside the included expense totals.'])
    summary.merge_cells(start_row=note_row+1,start_column=1,end_row=note_row+1,end_column=7)
    write(summary,note_row+1,['The Transactions worksheet lists all activity within this report’s business/year filters. Rules and reusable profile data remain in the app.'])
    for row in (note_row,note_row+1):summary.row_dimensions[row].height=28
    band(monthly,5,['Month','Business','Charges USD','Refunds USD','Net expense USD','Transactions'])
    for row,record in enumerate(data['monthly'],6):write(monthly,row,record[:2]+[money(v) for v in record[2:5]]+record[5:])
    write(monthly,6+len(data['monthly']),['TOTAL','',money(data['chargesCents']),money(data['refundsCents']),money(data['netCents']),data['includedCount']])
    for tab in workbook:
        header_rows={13,outside_row} if tab==summary else {5}
        bold_rows={1,8,total_row} if tab==summary else {1,6+len(data['monthly'])}
        for row in tab:
            for cell in row:
                if cell.row not in header_rows:
                    cell.font=Font(name='Arial',size=10,bold=cell.row in bold_rows)
                    cell.alignment=Alignment(vertical='center',horizontal='right' if isinstance(cell.value,(int,float,Decimal)) else 'left',wrap_text=True)
                if isinstance(cell.value,Decimal):cell.number_format=money_format
            if not tab.row_dimensions[row[0].row].height:tab.row_dimensions[row[0].row].height=23
        tab['A1'].font=Font(name='Arial',size=16,bold=True,color='384D38');tab.row_dimensions[1].height=30
        tab.print_options.horizontalCentered=True
        tab.print_area=f'A1:{get_column_letter(7 if tab==summary else 6)}{tab.max_row}'
    summary.column_dimensions['A'].width=30
    for row in range(outside_row+1,outside_row+len(data['outsideTotals'])+1):summary.row_dimensions[row].height=34
    monthly.freeze_panes='C6';monthly.print_title_rows='1:5'
    add_summary_transactions(workbook,rows,business_label,year)
    output=io.BytesIO();workbook.save(output);workbook.close()
    return output.getvalue()


def add_summary_transactions(workbook, rows, business_label, year):
    """One filterable detail row per transaction, including activity outside totals."""
    from math import ceil
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tab=workbook.create_sheet('Transactions')
    headers=TRANSACTION_HEADERS
    widths=[14,52,34,26,38,18,14,20,28,18,18,36,45,14,16,14,32,42,34,26,14,48,40,40,40]
    tab.merge_cells('A1:L1');tab.merge_cells('A2:L2');tab.merge_cells('A4:L4')
    tab['A1']='All transactions'
    tab['A2']=safe_text(business_label)
    tab['A3']='Reporting period';tab['B3']=safe_text(year or 'All recorded years')
    tab['D3']='Transactions';tab['E3']=len(rows)
    tab['A4']='Includes reviewed, pending, personal, excluded and transfer activity in the selected scope. Net expense USD contributes to the summary only when Included in totals is TRUE.'
    for column,header in enumerate(headers,1):
        cell=tab.cell(5,column,header)
        cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='384D38')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
        tab.column_dimensions[get_column_letter(column)].width=widths[column-1]
    def excel_date(value):
        try:return date.fromisoformat(value)
        except (ValueError,TypeError):return value or None
    for number,row in enumerate(sorted(rows,key=lambda r:(r['date'],r['id'])),6):
        values=transaction_values(row)
        values[0]=excel_date(values[0]);values[13]=excel_date(values[13])
        lines=1
        for column,value in enumerate(values,1):
            cell=tab.cell(number,column,safe_text(value) if isinstance(value,str) else value)
            cell.font=Font(name='Arial',size=10)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if number%2==0:cell.fill=PatternFill('solid',fgColor='F2F5EE')
            if column in (6,8):cell.number_format='#,##0.00;[Red](#,##0.00)'
            if column in (1,14):cell.number_format='mm/dd/yy'
            if isinstance(value,str):
                lines=max(lines,sum(max(1,ceil(len(line)/max(1,widths[column-1]-2))) for line in value.splitlines()))
        tab.row_dimensions[number].height=min(409,max(30,lines*15+6))
    for row in range(1,5):
        for cell in tab[row]:
            cell.font=Font(name='Arial',size=10)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
        tab.row_dimensions[row].height=24
    tab['A1'].font=Font(name='Arial',size=16,bold=True,color='384D38')
    tab.row_dimensions[1].height=30;tab.row_dimensions[4].height=32;tab.row_dimensions[5].height=34
    tab.freeze_panes='C6'
    tab.auto_filter.ref=f'A5:Y{max(5,tab.max_row)}'
    tab.sheet_view.showGridLines=False
    tab.print_title_rows='1:5'
    tab.sheet_properties.pageSetUpPr.fitToPage=True
    tab.page_setup.orientation='landscape';tab.page_setup.paperSize=tab.PAPERSIZE_A3
    tab.page_setup.fitToWidth=0;tab.page_setup.fitToHeight=0
    tab.print_area=f'A1:Y{tab.max_row}'


def pdf_summary(rows, business_label, year):
    """Printable summaries plus a complete transaction appendix."""
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, PageBreak

    rows=[dict(row) for row in rows]
    data=accountant_summary(rows,business_label,year)
    output=io.BytesIO()
    document=SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=30,rightMargin=30,topMargin=32,bottomMargin=30,
        title='Bookkeeper accountant report',author='Local Bookkeeper')
    body=ParagraphStyle('Body',fontName='Helvetica',fontSize=8,leading=10,spaceAfter=4)
    title=ParagraphStyle('Title',parent=body,fontSize=17,leading=22,textColor=colors.HexColor('#384D38'),spaceAfter=9)
    heading=ParagraphStyle('Heading',parent=body,fontSize=12,leading=16,spaceBefore=10,spaceAfter=7)
    white=ParagraphStyle('Header',parent=body,textColor=colors.white,fontName='Helvetica-Bold')
    def paragraph(value,style=body):return Paragraph(escape(str(value if value is not None else '')).replace('\n','<br/>'),style)
    def amount(cents):return format(Decimal(cents)/100,',.2f')
    def table(headers,records,widths):
        result=LongTable([[paragraph(h,white) for h in headers]]+[[paragraph(v) for v in record] for record in records],
            colWidths=widths,repeatRows=1,hAlign='LEFT',splitInRow=1)
        result.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#384D38')),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F2F5EE')]),
            ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),
            ('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        return result
    story=[paragraph('Accountant summary',title),paragraph(business_label),paragraph('Reporting period: '+(year or 'All recorded years')),
        paragraph('Included charges: $'+amount(data['chargesCents'])+'    Included refunds: $'+amount(data['refundsCents'])+'    Net expenses: $'+amount(data['netCents'])),
        paragraph(str(data['includedCount'])+' included transactions. '+str(data['pendingCount'])+' awaiting review.'),
        paragraph('Category totals',heading),
        table(['Year','Business','Category','Charges USD','Refunds USD','Net USD','Rows'],
            [r[:3]+[amount(v) for v in r[3:6]]+[str(r[6])] for r in data['categories']],[45,170,190,95,95,95,40]),
        paragraph('Activity outside expense totals',heading),
        table(['Reason','Transactions','Money out USD','Money in USD'],
            [[r[0],str(r[1]),amount(r[2]),amount(r[3])] for r in data['outsideTotals']],[350,100,140,140]),
        PageBreak(),paragraph('Monthly summary',title),paragraph(business_label+' - '+(year or 'All recorded years')),
        table(['Month','Business','Charges USD','Refunds USD','Net USD','Rows'],
            [r[:2]+[amount(v) for v in r[2:5]]+[str(r[5])] for r in data['monthly']],[70,240,120,120,120,60]),
        PageBreak(),paragraph('All transactions',title),paragraph(business_label+' - '+(year or 'All recorded years')),
        paragraph('Every transaction in the selected scope is listed. Only Net expense USD contributes to the summaries. Blank net amounts stay outside totals.')]
    details=[]
    for row in sorted(rows,key=lambda r:(r['date'],r['id'])):
        values=transaction_values(row)
        description=values[1]
        for label,index in [('Merchant',12),('Posted',13),('Credit decision',16),('Statement category',17),
                ('Source',18),('Sheet',19),('Row',20),('Notes',21),('ID',22),('Related ID',23),('Duplicate of',24)]:
            if values[index] not in (None,''):description+='\n'+label+': '+str(values[index])
        status=values[4]+'\n'+values[8]+'\n'+values[9]+'\n'+(values[11] or 'Included in totals')
        details.append([values[0],description,values[2]+'\n'+values[3],status,
            format(values[5],',.2f')+'\n'+values[6],'' if values[7] is None else format(values[7],',.2f')])
    story.append(table(['Date','Description / source references','Business / account','Category / status','Amount USD','Net expense USD'],
        details,[56,270,120,140,90,105]))
    if not rows:story.append(paragraph('No transactions in this business/year selection.'))
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#64705E'))
        canvas.drawString(30,16,'Local Bookkeeper - '+(year or 'All recorded years'))
        canvas.drawRightString(landscape(A4)[0]-30,16,'Page '+str(doc.page))
    document.build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()


def excel_report(rows, business_label, year, database=None, profile=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = ['Date','Description','Business','Source account','Category','Expense USD','Amount USD','Direction','Kind','Classification','Review status','Included in totals','Source document','Source sheet','Source row','Transaction ID','Duplicate of','Notes','Credit decision','Related transaction ID']
    def ledger_row(row):
        raw = json.loads(row['source_raw_json'])
        return [row['date'],safe_text(row['description']),safe_text(row['business_name']),safe_text(row['nickname'] + (' ···· '+row['last_four'] if row['last_four'] else '')),safe_text(row['category_name']),Decimal(expense_cents(row))/100 if included(row) else None,Decimal(row['amount_cents'])/100,row['direction'],row['transaction_kind'],row['classification'],row['status'],bool(included(row)),safe_text(row['filename']),safe_text(raw.get('sheet')),row['source_row'],row['id'],row['duplicate_of'],safe_text(row['notes']),CREDIT_LABELS.get(row['credit_treatment'],'') if row['direction']=='credit' or row['credit_treatment'] else '',row['related_transaction_id']]
    def sheet(name, titles, records, money_columns=()):
        tab = workbook.create_sheet(name)
        tab.append([safe_text(business_label) + ' · ' + str(year or 'All years')])
        tab.append(['Bookkeeping and tax preparation assistance; review tax treatment with your preparer.'])
        tab.append(titles)
        for record in records: tab.append(record)
        tab.freeze_panes = 'A4'
        tab.auto_filter.ref = f'A3:{get_column_letter(len(titles))}{max(3,tab.max_row)}'
        for cell in tab[3]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='384D38')
        for column in range(1,len(titles)+1):
            tab.column_dimensions[get_column_letter(column)].width = min(52,max(16,len(titles[column-1])+3))
        for number in money_columns:
            for cells in tab.iter_rows(min_row=4,min_col=number,max_col=number): cells[0].number_format = '#,##0.00;[Red](#,##0.00)'
    ledger = [r for r in rows if included(r)]
    sheet('Expense Ledger',headers,[ledger_row(r) for r in ledger],(6,7))
    monthly, categories = defaultdict(int), defaultdict(int)
    for row in ledger:
        monthly[row['date'][:7]] += expense_cents(row)
        categories[row['category_name']] += expense_cents(row)
    sheet('Monthly Summary',['Month','Expense USD'],[[k,Decimal(v)/100] for k,v in sorted(monthly.items())],(2,))
    sheet('Category Summary',['Category','Expense USD'],[[safe_text(k),Decimal(v)/100] for k,v in sorted(categories.items())],(2,))
    sheet('Tax Preparation Summary',['Category','Reviewed expense USD','Review guidance'],[[safe_text(k),Decimal(v)/100,'Confirm business purpose and tax treatment. Vehicle, meals, loan and owner activity need separate review.'] for k,v in sorted(categories.items())],(2,))
    sheet('Review Queue',headers,[ledger_row(r) for r in rows if r['status'] in {'needs_review','suggested'}],(6,7))
    sheet('Excluded Transactions',headers,[ledger_row(r) for r in rows if not included(r)],(6,7))
    sheet('Duplicate-Transfer Transactions',headers,[ledger_row(r) for r in rows if r['duplicate_of'] or r['transaction_kind']=='transfer' or r['classification']=='transfer'],(6,7))
    if database is not None: add_profile(workbook,database,profile or {})
    output = io.BytesIO(); workbook.save(output)
    return output.getvalue()
