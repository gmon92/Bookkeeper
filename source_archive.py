"""Immutable, database-resident source cells and transaction import baselines."""
import csv
import datetime as dt
import io
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile, BadZipFile

CATEGORY_HEADERS = ('category', 'transactioncategory', 'merchantcategory', 'expensecategory', 'categorydescription', 'originalcategory', 'statementcategory')


def category(raw, extraction=None):
    values = raw.get('values', [])
    column = (extraction or {}).get('columns', {}).get('sourceCategory')
    if column and 0 <= column['index'] < len(values):
        return str(values[column['index']] or '').strip()
    for header, value in zip(raw.get('headers', []), values):
        if re.sub(r'[^a-z0-9]', '', str(header).lower()) in CATEGORY_HEADERS:
            return str(value or '').strip()
    return ''


def category_groups(parsed):
    groups = {}
    for tx in parsed['rows']:
        name = category(tx['raw'], parsed['extraction'])
        group = groups.setdefault(name, {'category': name, 'count': 0, 'amountCents': 0})
        group['count'] += 1
        group['amountCents'] += tx['amount_cents']
    return sorted(groups.values(), key=lambda g: (-g['count'], g['category']))


def initialize(db, now, data_directory=None):
    db.executescript('''
    CREATE TABLE IF NOT EXISTS source_archives (
      document_id TEXT PRIMARY KEY REFERENCES documents(id), file_hash TEXT NOT NULL,
      coverage TEXT NOT NULL CHECK(coverage IN ('all_cells','imported_rows')),
      sheets_json TEXT NOT NULL, captured_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS source_rows (
      document_id TEXT NOT NULL REFERENCES source_archives(document_id), sheet TEXT NOT NULL,
      row_number INTEGER NOT NULL, cells_json TEXT NOT NULL,
      PRIMARY KEY(document_id,sheet,row_number)
    );
    CREATE TABLE IF NOT EXISTS original_transactions (
      transaction_id TEXT PRIMARY KEY REFERENCES transactions(id),
      document_id TEXT NOT NULL REFERENCES documents(id), sheet TEXT NOT NULL, source_row INTEGER,
      source_category TEXT NOT NULL, raw_json TEXT NOT NULL, normalized_json TEXT,
      captured_at TEXT NOT NULL
    );
    CREATE INDEX IF NOT EXISTS idx_original_category ON original_transactions(source_category);
    ''')
    # Existing source_raw_json is evidence, even if the working row was edited.
    # Do not invent its original normalized values from today's working values.
    for doc in db.execute('SELECT * FROM documents WHERE id NOT IN (SELECT document_id FROM source_archives)').fetchall():
        rows = db.execute('SELECT * FROM transactions WHERE document_id=?', (doc['id'],)).fetchall()
        sheets = {}
        for row in rows:
            raw = json.loads(row['source_raw_json'])
            sheet = str(raw.get('sheet') or '')
            sheets[sheet] = {'name': sheet, 'headers': raw.get('headers', [])}
        archive = None
        if data_directory:
            source = (data_directory / doc['stored_path']).resolve()
            if source.is_relative_to(data_directory.resolve()) and source.is_file() and source.stat().st_size <= 20 * 1024 * 1024:
                try:
                    content = source.read_bytes()
                    if hashlib.sha256(content).hexdigest() == doc['file_hash']:
                        archive = capture(content, doc['filename'])
                except (OSError, ValueError, BadZipFile, KeyError, TypeError):
                    pass
        if archive is not None:
            save_document(db, doc['id'], doc['file_hash'], archive, now)
        else:
            db.execute('INSERT INTO source_archives VALUES(?,?,?,?,?)',
                (doc['id'], doc['file_hash'], 'imported_rows', json.dumps(list(sheets.values())), now))
        for row in rows:
            raw = json.loads(row['source_raw_json'])
            sheet = str(raw.get('sheet') or '')
            if archive is None:
                db.execute('INSERT OR IGNORE INTO source_rows VALUES(?,?,?,?)',
                    (doc['id'], sheet, row['source_row'] or 0, json.dumps([{'type': 'text', 'value': v} for v in raw.get('values', [])])))
            save_transaction(db, row['id'], doc['id'], raw, row['source_row'], now)
    for table in ('source_archives', 'source_rows', 'original_transactions'):
        for operation in ('UPDATE', 'DELETE'):
            db.execute(f'''CREATE TRIGGER IF NOT EXISTS immutable_{table}_{operation.lower()}
              BEFORE {operation} ON {table} BEGIN SELECT RAISE(ABORT,'Original source records are read-only'); END''')
        match = {'source_archives':'document_id=NEW.document_id', 'source_rows':'document_id=NEW.document_id AND sheet=NEW.sheet AND row_number=NEW.row_number',
                 'original_transactions':'transaction_id=NEW.transaction_id'}[table]
        db.execute(f'''CREATE TRIGGER IF NOT EXISTS immutable_{table}_replace BEFORE INSERT ON {table}
          WHEN EXISTS(SELECT 1 FROM {table} WHERE {match})
          BEGIN SELECT RAISE(ABORT,'Original source records cannot be replaced'); END''')
    db.execute('''CREATE TRIGGER IF NOT EXISTS immutable_transaction_source BEFORE UPDATE OF
      source_raw_json,source_row,document_id,raw_amount_text,source_amount_cents,source_fingerprint ON transactions
      WHEN NEW.source_raw_json IS NOT OLD.source_raw_json OR NEW.source_row IS NOT OLD.source_row
        OR NEW.document_id IS NOT OLD.document_id OR NEW.raw_amount_text IS NOT OLD.raw_amount_text
        OR NEW.source_amount_cents IS NOT OLD.source_amount_cents OR NEW.source_fingerprint IS NOT OLD.source_fingerprint
      BEGIN SELECT RAISE(ABORT,'Original transaction evidence is read-only'); END''')


def capture(content, filename):
    """Preserve values/types/formulas, independent of the file and working rows.

    Empty XLSX cells are implicit in their row/column positions. Formatting,
    drawings and macros are file features, not part of this data archive.
    """
    result, sheets, total_cells, total_bytes = [], [], 0, 0

    def add(sheet, number, cells):
        nonlocal total_cells, total_bytes
        total_cells += len(cells)
        encoded = json.dumps(cells, ensure_ascii=False, allow_nan=False)
        total_bytes += len(encoded.encode('utf-8'))
        if len(result) >= 100000 or total_cells > 2000000 or total_bytes > 100 * 1024 * 1024:
            raise ValueError('Source data exceeds the database archive limit (100,000 rows, 2 million cells or 100 MB). Split the statement before importing.')
        result.append((sheet, number, encoded))

    if Path(filename).suffix.lower() == '.csv':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('cp1252')
        width = 0
        for number, values in enumerate(csv.reader(io.StringIO(text, newline='')), 1):
            width = max(width, len(values))
            add('', number, [{'type': 'text', 'value': value} for value in values])
        sheets.append({'name': '', 'rows': len(result), 'columns': width})
    else:
        from openpyxl import load_workbook
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds the source archive limit.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 16384:
                    raise ValueError('A worksheet exceeds the source archive dimensions. Export a smaller statement.')
                if (sheet.max_row or 0) * (sheet.max_column or 0) > 2000000:
                    raise ValueError('A worksheet exceeds the 2 million cell source archive limit. Export a smaller statement.')
                for number, row in enumerate(sheet.iter_rows(), 1):
                    cells = []
                    for cell in row:
                        value = cell.value
                        kind = cell.data_type
                        if isinstance(value, (dt.datetime, dt.date, dt.time)):
                            kind, value = type(value).__name__, value.isoformat()
                        elif isinstance(value, dt.timedelta):
                            kind, value = 'timedelta_seconds', value.total_seconds()
                        cells.append({'type': kind, 'value': value})
                    if any(c['value'] is not None for c in cells):
                        add(sheet.title, number, cells)
                sheets.append({'name': sheet.title, 'rows': sheet.max_row or 0, 'columns': sheet.max_column or 0})
        finally:
            workbook.close()
    return sheets, result


def save_document(db, document_id, file_hash, archive, now):
    sheets, rows = archive
    db.execute('INSERT INTO source_archives VALUES(?,?,?,?,?)',
        (document_id, file_hash, 'all_cells', json.dumps(sheets), now))
    db.executemany('INSERT INTO source_rows VALUES(?,?,?,?)',
        ((document_id, sheet, number, cells) for sheet, number, cells in rows))


def save_transaction(db, transaction_id, document_id, raw, row_number, now, normalized=None, extraction=None):
    db.execute('INSERT INTO original_transactions VALUES(?,?,?,?,?,?,?,?)',
        (transaction_id, document_id, str(raw.get('sheet') or ''), row_number, category(raw, extraction),
         json.dumps(raw), json.dumps(normalized) if normalized is not None else None, now))


def original(db, transaction_id):
    row = db.execute('SELECT * FROM original_transactions WHERE transaction_id=?', (transaction_id,)).fetchone()
    if row is None:
        return None
    return {'sourceCategory': row['source_category'], 'raw': json.loads(row['raw_json']),
            'normalized': json.loads(row['normalized_json']) if row['normalized_json'] else None,
            'capturedAt': row['captured_at']}


def document_rows(db, document_id, sheet, offset):
    archive = db.execute('SELECT * FROM source_archives WHERE document_id=?', (document_id,)).fetchone()
    if archive is None:
        raise ValueError('No database source archive exists for this document.')
    sheets = json.loads(archive['sheets_json'])
    if sheet is None:
        sheet = sheets[0]['name'] if sheets else ''
    if sheet not in {s['name'] for s in sheets}:
        raise ValueError('Choose a worksheet from this document.')
    total = db.execute('SELECT count(*) FROM source_rows WHERE document_id=? AND sheet=?', (document_id, sheet)).fetchone()[0]
    rows = db.execute('SELECT row_number,cells_json FROM source_rows WHERE document_id=? AND sheet=? ORDER BY row_number LIMIT 100 OFFSET ?', (document_id, sheet, offset))
    return {'coverage': archive['coverage'], 'fileHash': archive['file_hash'], 'capturedAt': archive['captured_at'],
            'sheets': sheets, 'sheet': sheet, 'offset': offset, 'total': total,
            'rows': [{'number': r['row_number'], 'cells': json.loads(r['cells_json'])} for r in rows]}
