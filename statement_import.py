"""Read structured statements without changing the source workbook."""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from import_evidence import finish

MAX_ROWS = 50000


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    return str(value)


def parse_xlsx(content, filename, csv_parser):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel support needs openpyxl. Run python -m pip install -r requirements.txt.") from exc
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds the 100 MB limit.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("This file is not a readable XLSX workbook.") from exc
    try:
        if '_Bookkeeper_Profile' in workbook.sheetnames:
            raise ValueError('This is a reusable Local Bookkeeper ledger. Use Open / switch profile → Restore export / backup to restore its transactions, rules and categorizations into new books.')
        candidates = []
        for sheet in workbook:
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number > 50:
                    break
                headers = [cell_text(v).strip() for v in values]
                names = {''.join(c for c in h.lower() if c.isalnum()) for h in headers}
                if names & {'date', 'transactiondate', 'purchasedate'} and names & {'description', 'merchant', 'payee'} and names & {'amount', 'transactionamount', 'debit', 'credit', 'withdrawal', 'deposit'}:
                    candidates.append((sheet, row_number, headers))
                    break
        if not candidates:
            raise ValueError("No transaction table found in the first 50 rows of any sheet. Expected Date, Description and Amount or Debit/Credit columns.")
        if len(candidates) > 1:
            raise ValueError("Multiple transaction tables found. Export the intended account table as a separate CSV or XLSX file.")
        sheet, header_row, headers = candidates[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + header_row:
            raise ValueError(f"Workbook exceeds the {MAX_ROWS:,} transaction-row limit.")
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(headers)
        source_rows = []
        blank_rows = 0
        for number, row in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if len(source_rows) >= MAX_ROWS:
                raise ValueError("Too many transaction rows.")
            if not any(cell.value is not None for cell in row):
                blank_rows += 1
                continue
            if any(cell.data_type == 'f' for cell in row):
                raise ValueError(f"Formula in transaction sheet {sheet.title}, row {number}. Export values only before importing.")
            writer.writerow([cell_text(cell.value) for cell in row])
            source_rows.append(number)
        parsed = csv_parser(output.getvalue().encode('utf-8'), filename)
        parsed['extraction'].update(headerRow=header_row,sheet=sheet.title,metadataRowsSkipped=header_row-1,
            blankRowsSkipped=blank_rows,otherSheets=[s.title for s in workbook if s.title!=sheet.title])
        for tx, source_row in zip(parsed['rows'], source_rows):
            tx['source_row'] = source_row
            tx['raw']['sheet'] = sheet.title
        parsed.update(format='xlsx', provider='Generic spreadsheet', parser_version='xlsx-v3', reconciliation=None)
        # A specific provider layout supplies a sign convention; the preview asks
        # the user to confirm it before any rows are committed.
        amex = sheet.title == 'Transaction Details' and {'Extended Details', 'Appears On Your Statement As', 'Reference', 'Category'} <= set(headers)
        if amex:
            parsed.update(provider='American Express', parser_version='amex-xlsx-v3')
            for tx in parsed['rows']:
                tx['direction'] = 'credit' if tx['source_amount_cents'] < 0 else 'debit'
                description = tx['description'].upper()
                payment = 'PAYMENT' in description and ('THANK' in description or 'RECEIVED' in description)
                tx['transaction_kind'] = 'transfer' if payment else 'expense' if tx['direction'] == 'debit' else 'unknown'
                fingerprint = '|'.join((tx['date'], tx['posted_date'] or '', str(tx['amount_cents']), tx['direction'], ''.join(c for c in tx['description'].lower() if c.isalnum())))
                tx['fingerprint'] = hashlib.sha256(fingerprint.encode()).hexdigest()
            if 'Transaction Summary' in workbook.sheetnames:
                expected = {}
                for row in workbook['Transaction Summary'].iter_rows(values_only=True):
                    label = cell_text(row[0]).strip().lower() if row else ''
                    if label in {'charges', 'payments & credits', 'total'} and len(row) > 1 and isinstance(row[1], (int, float, Decimal)):
                        expected[label] = int(Decimal(str(row[1])) * 100)
                actual = {'charges': sum(r['source_amount_cents'] for r in parsed['rows'] if r['source_amount_cents'] > 0), 'payments & credits': sum(r['source_amount_cents'] for r in parsed['rows'] if r['source_amount_cents'] < 0), 'total': sum(r['source_amount_cents'] for r in parsed['rows'])}
                if expected:
                    parsed['reconciliation'] = {'matched': all(actual[k] == v for k, v in expected.items()), 'expectedCents': expected, 'actualCents': actual, 'sheet': 'Transaction Summary'}
                    if not parsed['reconciliation']['matched']:
                        raise ValueError("Extracted amounts do not match the workbook summary. Import stopped; check the source file.")
        return finish(parsed)
    finally:
        workbook.close()


def parse_statement(content, filename, csv_parser):
    suffix = Path(filename).suffix.lower()
    if suffix == '.xlsx':
        return parse_xlsx(content, filename, csv_parser)
    if suffix == '.csv':
        parsed = csv_parser(content, filename)
        parsed.update(format='csv', provider='Generic CSV', parser_version='csv-v4', reconciliation=None)
        return finish(parsed)
    raise ValueError("Choose a CSV or XLSX statement. PDF and legacy XLS files are not supported yet.")
