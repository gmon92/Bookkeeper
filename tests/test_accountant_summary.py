"""Accountant exports reconcile summaries and filtered transaction details."""
import io
import csv
import unittest
from unittest.mock import patch
from decimal import Decimal

from openpyxl import load_workbook
import reports
import server
import test_credits


class AccountantSummaryTests(unittest.TestCase):
    def test_all_businesses_combines_new_business_and_omits_personal_details(self):
        with server.connect() as db:
            db.execute("INSERT INTO businesses(id,legal_name,display_name,created_at) VALUES('new-business','New business','New business',?)",(server.utc_now(),))
            db.execute("UPDATE transactions SET business_id='new-business' WHERE description='SUPPLIER PURCHASE'")
            db.execute("UPDATE transactions SET classification='personal',business_id=NULL,include_in_totals=0 WHERE description='PLUS MEMBERSHIP'")
            db.execute("INSERT INTO accounts(id,nickname) VALUES('personal-card','Personal card')")
        self.fixture.import_csv('personal-card',b'Date,Description,Debit,Credit\n10/01/2026,PERSONAL UNASSIGNED,500.00,\n10/02/2026,BUSINESS ON PERSONAL CARD,25.00,\n')
        with server.connect() as db:
            db.execute("UPDATE transactions SET classification='business',business_id='new-business' WHERE description='BUSINESS ON PERSONAL CARD'")
        workbook,_=self.export('?year=2026')
        self.addCleanup(workbook.close)
        detail=workbook['Transactions']
        records=[dict(zip([c.value for c in detail[5]],row)) for row in detail.iter_rows(min_row=6,values_only=True)]
        exported_ids={r['Transaction ID'] for r in records}
        self.assertNotIn(self.fixture.ids['PLUS MEMBERSHIP'],exported_ids)
        self.assertIn(self.fixture.ids['SUPPLIER PURCHASE'],exported_ids)
        self.assertTrue(all(r['Classification']!='personal' for r in records))
        self.assertNotIn('PERSONAL UNASSIGNED',{r['Description'] for r in records})
        self.assertIn('BUSINESS ON PERSONAL CARD',{r['Description'] for r in records})
        self.assertEqual(workbook['Accountant Summary']['B8'].value,60)
        self.assertIn('New business',{c.value for row in workbook['Accountant Summary'] for c in row})
        self.handler.path='/api/export.csv?year=2026&view=transactions'
        csv_rows=list(csv.DictReader(io.StringIO(self.handler.export_csv()[0].decode('utf-8-sig'))))
        self.assertEqual({r['Transaction ID'] for r in csv_rows},exported_ids)
        self.handler.path='/api/export.pdf?year=2026'
        with patch('server.pdf_summary',return_value=b'%PDF-example') as renderer:
            self.handler.export_pdf()
            self.assertEqual({r['id'] for r in renderer.call_args.args[0]},exported_ids)

    def setUp(self):
        self.addCleanup(setattr,server,'PROFILE_PATH',server.PROFILE_PATH)
        self.fixture=test_credits.CreditReviewTests()
        self.fixture.setUp();self.addCleanup(self.fixture.tearDown)
        self.handler=self.fixture.handler
        self.handler.download=lambda data,name,kind:(data,name,kind)

    def export(self,query=''):
        self.handler.path='/api/export.xlsx'+query
        content,name,_=self.handler.export_excel()
        return load_workbook(io.BytesIO(content),data_only=True),name

    def test_default_export_is_a_summary_and_refunds_reconcile(self):
        self.fixture.apply({'ids':[self.fixture.ids['SUPPLIER REFUND']],'treatment':'refund','refundUsage':'business',
            'businessId':self.fixture.business,'category':'Office supplies','includeInTotals':True})
        workbook,name=self.export()
        self.addCleanup(workbook.close)
        self.assertEqual(name,'bookkeeper-accountant-summary.xlsx')
        self.assertEqual(workbook.sheetnames,['Accountant Summary','Monthly Summary','Transactions'])
        summary=workbook['Accountant Summary'];monthly=workbook['Monthly Summary']
        self.assertEqual((summary['B6'].value,summary['B7'].value,summary['B8'].value),(72.95,20,52.95))
        self.assertEqual((summary['D14'].value,summary['E14'].value,summary['F14'].value,summary['G14'].value),(72.95,20,52.95,3))
        self.assertEqual((monthly['C6'].value,monthly['D6'].value,monthly['E6'].value),(72.95,20,52.95))
        all_values={str(cell.value) for tab in workbook for row in tab for cell in row}
        self.assertIn('SUPPLIER PURCHASE',all_values)
        self.assertTrue(all(tx in all_values for tx in self.fixture.ids.values()))
        self.assertNotIn(reports.PROFILE_SHEET,workbook.sheetnames)
        self.assertTrue(all(tab.sheet_state=='visible' for tab in workbook))
        self.assertEqual(summary['B8'].data_type,'n')
        self.assertIn('A1:F1',{str(r) for r in monthly.merged_cells.ranges})
        self.assertIn('A2:F2',{str(r) for r in monthly.merged_cells.ranges})
        detail=workbook['Transactions']
        records=[dict(zip([c.value for c in detail[5]],row)) for row in detail.iter_rows(min_row=6,values_only=True)]
        self.assertEqual(len(records),7)
        self.assertEqual(len({r['Transaction ID'] for r in records}),7)
        self.assertEqual(sum(Decimal(str(r['Net expense USD'] or 0)) for r in records),Decimal('52.95'))
        refund=next(r for r in records if r['Description']=='SUPPLIER REFUND')
        self.assertEqual((refund['Amount USD'],refund['Net expense USD'],refund['Direction']),(20,-20,'credit'))
        self.assertTrue(refund['Included in totals'])
        pending=next(r for r in records if r['Description']=='CARD WITHDRAWAL')
        self.assertEqual((pending['Review status'],pending['Outside totals reason']),('Awaiting review','Awaiting review'))
        self.assertIsNone(pending['Net expense USD'])
        self.assertEqual(detail.freeze_panes,'C6')
        self.assertEqual(detail.auto_filter.ref,'A5:Y12')
        self.assertEqual(detail['A6'].data_type,'d')
        self.assertEqual(detail['F6'].data_type,'n')

    def test_business_and_year_filters_do_not_leak_other_business_on_same_account(self):
        with server.connect() as db:
            db.execute("INSERT INTO businesses(id,legal_name,display_name,created_at) VALUES('other','Other books','Other books',?)",(server.utc_now(),))
            db.execute("UPDATE transactions SET business_id='other' WHERE description='PLUS MEMBERSHIP'")
        self.fixture.import_csv('card',b'Date,Description,Debit,Credit\n10/01/2025,PREVIOUS YEAR PRIVATE,99.00,\n')
        workbook,_=self.export('?businessId='+self.fixture.business+'&year=2026')
        self.addCleanup(workbook.close)
        summary=workbook['Accountant Summary']
        self.assertEqual(summary['B8'].value,60)
        self.assertEqual(summary['B3'].value,'2026')
        self.assertNotIn('Other books',{cell.value for tab in workbook for row in tab for cell in row})
        detail=workbook['Transactions']
        values={cell.value for row in detail for cell in row}
        self.assertNotIn('PLUS MEMBERSHIP',values)
        self.assertNotIn(self.fixture.ids['PLUS MEMBERSHIP'],values)
        self.assertNotIn('PREVIOUS YEAR PRIVATE',values)
        self.assertIn('SUPPLIER PURCHASE',values)

    def test_outside_groups_partition_all_rows_without_double_counting(self):
        with server.connect() as db:
            db.execute("UPDATE transactions SET classification='personal',business_id=NULL,excluded=1,include_in_totals=0 WHERE description='PLUS MEMBERSHIP'")
            rows=db.execute('SELECT t.*,b.display_name AS business_name FROM transactions t LEFT JOIN businesses b ON b.id=t.business_id').fetchall()
            result=reports.accountant_summary(rows,'All businesses','')
            self.assertEqual(result['includedCount'],1)
            self.assertEqual(result['netCents'],6000)
            self.assertEqual(result['includedCount']+sum(r[1] for r in result['outsideTotals']),len(rows))
            personal=next(r for r in result['outsideTotals'] if r[0]=='Personal activity')
            self.assertEqual(personal,['Personal activity',1,1295,0])
            self.assertEqual(sum(r[2] for r in result['outsideTotals'])+result['chargesCents'],sum(r['amount_cents'] for r in rows if r['direction']=='debit'))
            self.assertEqual(sum(r[3] for r in result['outsideTotals'])+result['refundsCents'],sum(r['amount_cents'] for r in rows if r['direction']=='credit'))

    def test_summary_rejects_embedded_profile_and_exports_empty_period(self):
        self.handler.path='/api/export.xlsx?portable=1'
        with self.assertRaisesRegex(ValueError,'never include a reusable profile'):self.handler.export_excel()
        workbook,_=self.export('?year=2020')
        self.addCleanup(workbook.close)
        self.assertEqual(workbook['Accountant Summary']['B8'].value,0)
        self.assertEqual(workbook['Monthly Summary']['A6'].value,'TOTAL')
        self.assertEqual(workbook['Transactions'].max_row,5)
        self.assertEqual(workbook['Transactions'].auto_filter.ref,'A5:Y5')

    def test_transaction_details_preserve_outside_activity_and_escape_source_text(self):
        with server.connect() as db:
            db.execute("UPDATE transactions SET description='=2+2',merchant='@merchant',notes='+private note',classification='business',include_in_totals=0,excluded=1 WHERE description='PLUS MEMBERSHIP'")
            before=[tuple(r) for r in db.execute('SELECT * FROM transactions ORDER BY id')]
        self.fixture.apply({'ids':[self.fixture.ids['Platinum Walmart+ Credit']],'treatment':'statement_credit'})
        workbook,_=self.export()
        self.addCleanup(workbook.close)
        detail=workbook['Transactions']
        records=[dict(zip([c.value for c in detail[5]],row)) for row in detail.iter_rows(min_row=6,values_only=True)]
        personal=next(r for r in records if r['Transaction ID']==self.fixture.ids['PLUS MEMBERSHIP'])
        self.assertEqual((personal['Description'],personal['Merchant'],personal['Notes']),("'=2+2","'@merchant","'+private note"))
        self.assertEqual((personal['Classification'],personal['Outside totals reason'],personal['Excluded']),('business','Explicitly excluded activity',True))
        self.assertIsNone(personal['Net expense USD'])
        # This benefit is classified Personal by its saved credit decision.
        self.assertNotIn(self.fixture.ids['Platinum Walmart+ Credit'],{r['Transaction ID'] for r in records})
        self.assertTrue(all(cell.data_type!='f' for row in detail for cell in row))
        self.assertEqual(personal['Statement category'],'Benefits')
        self.assertEqual(personal['Source document'],'example.csv')
        self.assertIsNone(personal['Credit decision'])
        self.assertEqual(personal['Date'].strftime('%Y-%m-%d'),'2026-10-01')
        with server.connect() as db:
            current=[tuple(r) for r in db.execute('SELECT * FROM transactions ORDER BY id')]
        # The explicit credit decision changed its row; exporting did not touch any other row.
        benefit_id=self.fixture.ids['Platinum Walmart+ Credit']
        self.assertEqual([r for r in current if r[0]!=benefit_id],[r for r in before if r[0]!=benefit_id])

    def test_older_personal_category_rows_never_inflate_expense_totals(self):
        with server.connect() as db:
            db.execute("UPDATE transactions SET category_name='Personal / non-business' WHERE description='PLUS MEMBERSHIP'")
            row=db.execute("SELECT * FROM transactions WHERE description='PLUS MEMBERSHIP'").fetchone()
            self.assertFalse(reports.included(row))
            self.assertEqual(row['include_in_totals'],1)  # Reporting does not rewrite saved decisions.
        workbook,_=self.export()
        self.addCleanup(workbook.close)
        self.assertEqual(workbook['Accountant Summary']['B8'].value,60)

    def test_all_transactions_csv_matches_excel_details_and_escapes_formulas(self):
        with server.connect() as db:
            db.execute("UPDATE transactions SET description='=2+2',notes='+notes' WHERE description='PLUS MEMBERSHIP'")
        self.handler.path='/api/export.csv?view=transactions'
        content,name,kind=self.handler.export_csv()
        self.assertEqual((name,kind),('bookkeeper-transactions.csv','text/csv; charset=utf-8'))
        records=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
        workbook,_=self.export();self.addCleanup(workbook.close)
        sheet=workbook['Transactions'];headers=[c.value for c in sheet[5]]
        excel=[dict(zip(headers,row)) for row in sheet.iter_rows(min_row=6,values_only=True)]
        self.assertEqual({r['Transaction ID'] for r in records},{r['Transaction ID'] for r in excel})
        self.assertEqual(len(records),7)
        self.assertEqual(sum(Decimal(r['Net expense USD'] or 0) for r in records),Decimal('72.95'))
        personal=next(r for r in records if r['Transaction ID']==self.fixture.ids['PLUS MEMBERSHIP'])
        self.assertEqual((personal['Description'],personal['Notes']),("'=2+2","'+notes"))

    def test_csv_and_pdf_share_excel_business_year_filters_and_handle_empty_reports(self):
        with server.connect() as db:
            db.execute("INSERT INTO businesses(id,legal_name,display_name,created_at) VALUES('other','Other books','Other books',?)",(server.utc_now(),))
            db.execute("UPDATE transactions SET business_id='other' WHERE description='PLUS MEMBERSHIP'")
        self.fixture.import_csv('card',b'Date,Description,Debit,Credit\n10/01/2025,PREVIOUS YEAR PRIVATE,99.00,\n')
        suffix='?businessId='+self.fixture.business+'&year=2026'
        self.handler.path='/api/export.csv'+suffix+'&view=transactions'
        records=list(csv.DictReader(io.StringIO(self.handler.export_csv()[0].decode('utf-8-sig'))))
        self.assertNotIn('PLUS MEMBERSHIP',{r['Description'] for r in records})
        self.assertNotIn('PREVIOUS YEAR PRIVATE',{r['Description'] for r in records})
        with patch.object(server,'pdf_summary',wraps=reports.pdf_summary) as renderer:
            self.handler.path='/api/export.pdf'+suffix
            content,name,kind=self.handler.export_pdf()
            self.assertTrue(content.startswith(b'%PDF-'))
            self.assertEqual((name,kind),('bookkeeper-accountant-report.pdf','application/pdf'))
            exported_rows,label,year=renderer.call_args.args
            self.assertEqual({r['id'] for r in exported_rows},{r['Transaction ID'] for r in records})
            self.assertEqual((label,year),('Example Business','2026'))
        self.handler.path='/api/export.pdf?year=2020'
        self.assertTrue(self.handler.export_pdf()[0].startswith(b'%PDF-'))
        for path in ['/api/export.pdf?businessId=missing','/api/export.csv?view=transactions&year=20255','/api/export.csv?view=invalid']:
            self.handler.path=path
            with self.assertRaises(ValueError):
                self.handler.export_pdf() if '.pdf' in path else self.handler.export_csv()

    def test_pdf_wraps_long_notes_and_literal_markup(self):
        with server.connect() as db:
            db.execute("UPDATE transactions SET description='<invoice> & payment',notes=? WHERE description='SUPPLIER PURCHASE'",('Long note & detail\n'*400,))
        self.handler.path='/api/export.pdf'
        self.assertTrue(self.handler.export_pdf()[0].startswith(b'%PDF-'))
