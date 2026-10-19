"""Financial behavior: exclusions, paired payments, refunds, rules and undo."""
import hashlib
import base64
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import time
import unittest

import credits
import reports
import server
import source_archive


class CreditReviewTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='bookkeeper-credit-tests-')
        self.root = Path(self.temporary.name)
        self.previous = server.DATA_DIR,server.DOCS_DIR,server.DB_PATH,server.WORKSPACE_READY,server.WORKSPACE_ID
        server.activate_workspace(self.root,'Example Business')
        self.handler = server.Handler.__new__(server.Handler)
        self.handler.send_json = lambda status,body:(status,body)
        with server.connect() as db:
            self.business = db.execute('SELECT id FROM businesses').fetchone()[0]
            self.category = db.execute("SELECT id,name FROM categories WHERE name='Office supplies'").fetchone()
            db.execute("INSERT INTO accounts(id,nickname,business_id) VALUES('card','Example card',?)",(self.business,))
            db.execute("INSERT INTO accounts(id,nickname,business_id) VALUES('bank','Example bank',?)",(self.business,))
        self.import_csv('card',b'Date,Description,Debit,Credit,Category\n10/01/2026,PLUS MEMBERSHIP,12.95,,Benefits\n10/02/2026,Platinum Walmart+ Credit,,12.95,Benefits\n10/03/2026,MOBILE PAYMENT - THANK YOU,,100.00,Payments\n10/04/2026,SUPPLIER PURCHASE,60.00,,Supplies\n10/05/2026,SUPPLIER REFUND,,20.00,Supplies\n10/06/2026,MOBILE PAYMENT - THANK YOU,,110.00,Payments\n')
        self.import_csv('bank',b'Date,Description,Debit,Credit\n10/03/2026,CARD WITHDRAWAL,100.00,\n')
        with server.connect() as db:
            self.ids = {r['description']:r['id'] for r in db.execute('SELECT * FROM transactions')}
            # Membership and purchase have already been explicitly included.
            db.execute("UPDATE transactions SET classification='business',business_id=?,category_id=?,category_name=?,status='approved',include_in_totals=1 WHERE direction='debit' AND account_id='card'",(self.business,self.category['id'],self.category['name']))

    def tearDown(self):
        server.DATA_DIR,server.DOCS_DIR,server.DB_PATH,server.WORKSPACE_READY,server.WORKSPACE_ID = self.previous
        self.temporary.cleanup()

    def import_csv(self,account,content):
        parsed = server.parse_statement(content,'example.csv',server.parse_csv)
        token = hashlib.sha256(content).hexdigest()
        with server.connect() as db:
            db.execute('INSERT INTO staged_imports VALUES(?,?,?,?,?)',(token,content,json.dumps(parsed),token,time.time()))
        self.handler.read_json=lambda:{'token':token,'accountId':account,'confirmExtraction':True}
        status,_=self.handler.commit_import()
        self.assertEqual(status,201)

    def row(self,name):
        with server.connect() as db:return dict(db.execute('SELECT * FROM transactions WHERE id=?',(self.ids[name],)).fetchone())

    def apply(self,body):
        with server.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            plan=credits.plan(db,body,self.handler.validate_transaction)[0]
            return credits.commit(db,body|{'revision':plan['revision']},self.handler,server.utc_now())

    def expense_total(self):
        with server.connect() as db:return sum(reports.expense_cents(r) for r in db.execute('SELECT * FROM transactions') if reports.included(r))

    def test_statement_benefit_excludes_only_credit_and_keeps_original_charge(self):
        before=self.row('PLUS MEMBERSHIP')
        with server.connect() as db: original=source_archive.original(db,self.ids['Platinum Walmart+ Credit'])
        self.apply({'ids':[self.ids['Platinum Walmart+ Credit']],'treatment':'statement_credit'})
        credit=self.row('Platinum Walmart+ Credit')
        self.assertEqual((credit['excluded'],credit['include_in_totals'],credit['classification'],credit['business_id']),(1,0,'personal',None))
        self.assertEqual(self.row('PLUS MEMBERSHIP'),before)
        self.assertEqual(self.expense_total(),7295)
        with server.connect() as db:
            self.assertEqual(source_archive.original(db,credit['id']),original)
            with self.assertRaises(sqlite3.IntegrityError):db.execute('UPDATE transactions SET include_in_totals=1 WHERE id=?',(credit['id'],))

    def test_linked_payment_handles_both_sides_and_one_undo_restores_them(self):
        payment=self.ids['MOBILE PAYMENT - THANK YOU']
        with server.connect() as db:
            payment=db.execute("SELECT id FROM transactions WHERE account_id='card' AND direction='credit' AND amount_cents=10000").fetchone()[0]
            candidates=credits.candidates(db,payment,'payment')
        self.assertEqual([r['id'] for r in candidates],[self.ids['CARD WITHDRAWAL']])
        before=self.row('CARD WITHDRAWAL')
        result=self.apply({'ids':[payment],'treatment':'payment','relatedId':before['id']})
        self.assertEqual(result['affectedCount'],2)
        with server.connect() as db:
            pair=[dict(db.execute('SELECT * FROM transactions WHERE id=?',(tx,)).fetchone()) for tx in (payment,before['id'])]
            self.assertTrue(all(r['classification']=='transfer' and not reports.included(r) for r in pair))
            self.assertEqual([r['related_transaction_id'] for r in pair],[before['id'],payment])
        self.handler.read_json=lambda:{'actionId':result['actionId']}
        status,undo=self.handler.undo_review()
        self.assertEqual(status,200)
        self.assertEqual(len(undo['restored']),2)
        self.assertEqual(self.row('CARD WITHDRAWAL')['classification'],before['classification'])
        self.assertIsNone(self.row('CARD WITHDRAWAL')['related_transaction_id'])

    def test_business_refund_needs_explicit_inclusion_and_never_edits_purchase(self):
        purchase=self.row('SUPPLIER PURCHASE')
        body={'ids':[self.ids['SUPPLIER REFUND']],'treatment':'refund','refundUsage':'business','businessId':self.business,'category':self.category['name'],'relatedId':purchase['id']}
        self.apply(body)
        self.assertEqual(self.expense_total(),7295)
        self.apply(body|{'includeInTotals':True})
        self.assertEqual(self.expense_total(),5295)
        self.assertEqual(self.row('SUPPLIER PURCHASE'),purchase)
        self.apply({'ids':body['ids'],'treatment':'refund','refundUsage':'outside'})
        self.assertEqual(self.expense_total(),7295)
        self.assertEqual(self.row('SUPPLIER REFUND')['excluded'],1)

    def test_mixed_selection_invalid_links_and_stale_previews_save_nothing(self):
        payment=self.ids['Platinum Walmart+ Credit']
        with server.connect() as db:
            with self.assertRaises(ValueError):credits.plan(db,{'ids':[payment,self.ids['PLUS MEMBERSHIP']],'treatment':'statement_credit'},self.handler.validate_transaction)
            with self.assertRaises(ValueError):credits.plan(db,{'ids':[payment],'treatment':'payment','relatedId':self.ids['PLUS MEMBERSHIP']},self.handler.validate_transaction)
            with self.assertRaises(ValueError):credits.plan(db,{'ids':[payment],'treatment':'statement_credit','includeInTotals':True},self.handler.validate_transaction)
            body={'ids':[payment],'treatment':'statement_credit'}
            plan=credits.plan(db,body,self.handler.validate_transaction)[0]
            db.execute("UPDATE transactions SET notes='Edited in another view' WHERE id=?",(payment,))
            with self.assertRaisesRegex(ValueError,'changed'):credits.commit(db,body|{'revision':plan['revision']},self.handler,server.utc_now())
            self.assertEqual(db.execute('SELECT count(*) FROM review_actions').fetchone()[0],0)
        self.assertEqual(self.row('Platinum Walmart+ Credit')['credit_treatment'],'')

    def test_saved_rules_suggest_on_credits_only_and_conflicts_need_a_decision(self):
        with server.connect() as db: ids=[r[0] for r in db.execute("SELECT id FROM transactions WHERE description LIKE 'MOBILE PAYMENT%'")]
        result=self.apply({'ids':ids,'treatment':'payment','rememberRule':True,'rulePattern':'MOBILE PAYMENT','ruleAccountId':'card'})
        self.handler.read_json=lambda:{'actionId':result['actionId']}
        self.handler.undo_review()
        with server.connect() as db:
            self.assertEqual(len(credits.rules(db)),1)  # Undo preserves a separately saved suggestion rule.
            suggestions=credits.suggestions(db)
            self.assertTrue(all(suggestions[tx]['treatment']=='payment' for tx in ids))
            self.assertNotIn(self.ids['PLUS MEMBERSHIP'],suggestions)
            self.assertTrue(all(db.execute('SELECT status FROM transactions WHERE id=?',(tx,)).fetchone()[0]=='needs_review' for tx in ids))
            db.execute("INSERT INTO credit_rules(id,pattern,account_id,treatment,created_at) VALUES('conflict','MOBILE PAYMENT','card','excluded_credit',?)",(server.utc_now(),))
            self.assertTrue(credits.suggestions(db)[ids[0]]['conflict'])

    def test_unclear_credit_returns_to_review_outside_totals(self):
        credit=self.ids['SUPPLIER REFUND']
        self.apply({'ids':[credit],'treatment':'refund','refundUsage':'business','businessId':self.business,'category':self.category['name'],'includeInTotals':True})
        self.apply({'ids':[credit],'treatment':''})
        row=self.row('SUPPLIER REFUND')
        self.assertEqual((row['status'],row['classification'],row['include_in_totals'],row['credit_treatment']),('needs_review','unknown',0,''))
        self.assertEqual(self.expense_total(),7295)

    def test_http_dispatch_preview_and_apply(self):
        body={'ids':[self.ids['Platinum Walmart+ Credit']],'treatment':'statement_credit'}
        self.handler.local_request=lambda:True
        self.handler.headers={'X-Bookkeeper-Workspace':server.WORKSPACE_ID}
        self.handler.path='/api/credits/preview'
        self.handler.read_json=lambda:body
        status,plan=self.handler.post_request()
        self.assertEqual(status,200)
        self.assertEqual(plan['selectedCount'],1)
        self.assertEqual(self.row('Platinum Walmart+ Credit')['credit_treatment'],'')
        body['revision']=plan['revision']
        self.handler.path='/api/credits/apply'
        status,result=self.handler.post_request()
        self.assertEqual(status,200)
        self.assertEqual(result['affectedCount'],1)
        self.assertEqual(self.row('Platinum Walmart+ Credit')['credit_treatment'],'statement_credit')
        self.handler.path='/api/imports/preview'
        self.handler.read_json=lambda:{'filename':'another.csv','contentBase64':base64.b64encode(b'Date,Description,Debit,Credit\n10/10/2026,EXAMPLE CREDIT,,12.95\n').decode()}
        status,preview=self.handler.post_request()
        self.assertEqual(status,200)
        self.assertEqual(preview['creditCents'],1295)

    def test_unlinking_a_payment_releases_the_bank_withdrawal_for_ordinary_review(self):
        with server.connect() as db: payment=db.execute("SELECT id FROM transactions WHERE account_id='card' AND direction='credit' AND amount_cents=10000").fetchone()[0]
        bank=self.ids['CARD WITHDRAWAL']
        self.apply({'ids':[payment],'treatment':'payment','relatedId':bank})
        self.apply({'ids':[payment],'treatment':'payment','relatedId':None})
        released=self.row('CARD WITHDRAWAL')
        self.assertEqual(released['credit_treatment'],'')
        self.assertIsNone(released['related_transaction_id'])
        released.update(transaction_kind='expense',classification='business',business_id=self.business,category_id=self.category['id'],category_name=self.category['name'])
        with server.connect() as db:self.handler.validate_transaction(db,released)

    def test_excel_report_keeps_benefit_traceability_outside_expense_ledger(self):
        from openpyxl import load_workbook
        self.apply({'ids':[self.ids['Platinum Walmart+ Credit']],'treatment':'statement_credit'})
        with server.connect() as db:
            rows=db.execute('''SELECT t.*,a.nickname,a.last_four,d.filename,b.display_name AS business_name FROM transactions t
              JOIN accounts a ON a.id=t.account_id LEFT JOIN documents d ON d.id=t.document_id LEFT JOIN businesses b ON b.id=t.business_id''').fetchall()
        workbook=load_workbook(io.BytesIO(reports.excel_report(rows,'Example Business','2026')),read_only=True,data_only=True)
        try:
            ledger=list(workbook['Expense Ledger'].iter_rows(min_row=4,values_only=True))
            self.assertEqual({r[1] for r in ledger},{'PLUS MEMBERSHIP','SUPPLIER PURCHASE'})
            self.assertEqual(sum(r[5] for r in ledger),72.95)
            headers=next(workbook['Excluded Transactions'].iter_rows(min_row=3,max_row=3,values_only=True))
            benefit=next(r for r in workbook['Excluded Transactions'].iter_rows(min_row=4,values_only=True) if r[1]=='Platinum Walmart+ Credit')
            self.assertEqual(benefit[headers.index('Credit decision')],'Statement credit / benefit')
            self.assertFalse(benefit[headers.index('Included in totals')])
        finally:workbook.close()

    def test_server_does_not_allow_another_listener_to_serve_stale_code(self):
        with server.LocalHTTPServer(('127.0.0.1',0),server.Handler) as active:
            with self.assertRaises(OSError):server.LocalHTTPServer(active.server_address,server.Handler)


if __name__=='__main__':unittest.main()
