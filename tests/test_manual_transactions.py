"""Manual entries preserve original input, totals safeguards, and duplicate matching."""
import copy
import csv
import io
import json
import sqlite3
import unittest
import uuid

import automation
import reports
import server
import source_archive
import test_credits


class ManualTransactionTests(unittest.TestCase):
    def setUp(self):
        self.addCleanup(setattr,server,'PROFILE_PATH',server.PROFILE_PATH)
        self.fixture=test_credits.CreditReviewTests();self.fixture.setUp()
        self.addCleanup(self.fixture.tearDown)
        self.handler=self.fixture.handler;self.handler.local_request=lambda:True
        self.handler.headers={'X-Bookkeeper-Workspace':server.WORKSPACE_ID}

    def body(self,**changes):
        return {'requestId':str(uuid.uuid4()),'date':'2026-10-22','description':'MANUAL SHOP',
            'amountCents':1525,'accountId':'card','direction':'debit','kind':'expense',
            'classification':'unknown','category':'Uncategorized / review','notes':'Receipt entered by hand'}|changes

    def request(self,path,body):
        self.handler.path=path;self.handler.read_json=lambda:body
        return self.handler.post_request()

    def row(self,identity):
        with server.connect() as db:return dict(db.execute('SELECT * FROM transactions WHERE id=?',(identity,)).fetchone())

    def test_manual_entry_without_a_document_starts_outside_totals_and_keeps_original(self):
        with server.connect() as db:
            before=db.execute('SELECT count(*) FROM documents').fetchone()[0]
            automation.save(db,{'type':'expense','patterns':['MANUAL SHOP'],'accountId':None,
                'businessId':self.fixture.business,'category':'Meals','mode':'review','includeInTotals':True},server.utc_now())
        body=self.body();status,result=self.request('/api/transactions',body)
        self.assertEqual(status,201,result);row=self.row(result['id'])
        self.assertEqual((row['status'],row['classification'],row['include_in_totals']),('needs_review','unknown',0))
        self.assertIsNone(row['document_id'])
        with server.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM documents').fetchone()[0],before)
            original=source_archive.original(db,result['id'])
            self.assertEqual((original['origin'],original['normalized']['amountCents']),('manual',1525))
            self.assertEqual(source_archive.category(original['raw']),'')
            self.assertTrue(db.execute('SELECT 1 FROM transaction_history WHERE transaction_id=? AND action=?',(result['id'],'manual_entry')).fetchone())
            with self.assertRaises(sqlite3.IntegrityError):
                db.execute('UPDATE transactions SET source_raw_json=? WHERE id=?',('{}',result['id']))
        self.handler.path='/api/transactions'
        status,rows=self.handler.get_request()
        saved=next(r for r in rows if r['id']==result['id'])
        self.assertEqual((saved['detail'],saved['sourceCategory']),('Manual entry',''))

    def test_explicit_business_expense_and_refund_contribute_correctly(self):
        common={'classification':'business','businessId':self.fixture.business,'category':'Office supplies',
            'includeInTotals':True,'status':'approved'}
        _,expense=self.request('/api/transactions',self.body(**common))
        status,refund=self.request('/api/transactions',self.body(**common,description='MANUAL REFUND',direction='credit',kind='refund',amountCents=500))
        self.assertEqual(status,201,refund)
        self.assertEqual(reports.expense_cents(self.row(expense['id'])),1525)
        refund_row=self.row(refund['id'])
        self.assertEqual((refund_row['credit_treatment'],reports.expense_cents(refund_row)),('refund',-500))

    def test_personal_and_transfer_entries_stay_outside_business_totals(self):
        for changes in ({'classification':'personal','category':'Personal / non-business'},
                {'classification':'transfer','kind':'transfer'}):
            status,result=self.request('/api/transactions',self.body(**changes))
            self.assertEqual(status,201,result);self.assertFalse(reports.included(self.row(result['id'])))
            status,_=self.request('/api/transactions',self.body(**changes,includeInTotals=True,status='approved'))
            self.assertEqual(status,400)

    def test_invalid_fields_save_nothing(self):
        invalid=[{'amountCents':0},{'amountCents':True},{'amountCents':1.25},{'amountCents':-1},
            {'date':'2026-02-30'},{'description':'  '},{'accountId':'missing'},{'accountId':[]},
            {'classification':'business'},{'classification':[]},{'status':[]},{'kind':[]},
            {'direction':'credit'},{'direction':'unknown'},{'includeInTotals':'true'},
            {'classification':'business','businessId':self.fixture.business,'category':'Personal / non-business'},
            {'classification':'business','businessId':self.fixture.business,'category':'Not a category'},
            {'requestId':'not-an-id'},{'rememberRule':True}]
        with server.connect() as db:before=db.execute('SELECT count(*) FROM transactions').fetchone()[0]
        for changes in invalid:
            with self.subTest(changes=changes):
                status,result=self.request('/api/transactions',self.body(**changes))
                self.assertEqual(status,400,result)
        with server.connect() as db:self.assertEqual(db.execute('SELECT count(*) FROM transactions').fetchone()[0],before)

    def test_repeated_save_returns_same_transaction_but_changed_request_is_rejected(self):
        body=self.body();_,first=self.request('/api/transactions',body)
        status,second=self.request('/api/transactions',body)
        self.assertEqual((status,second['id']),(201,first['id']))
        status,result=self.request('/api/transactions',body|{'amountCents':2000})
        self.assertEqual(status,400)
        self.assertIn('already saved',result['error'])
        with server.connect() as db:self.assertEqual(db.execute("SELECT count(*) FROM transactions WHERE description='MANUAL SHOP'").fetchone()[0],1)

    def test_later_edit_keeps_original_values_and_creation_history(self):
        _,saved=self.request('/api/transactions',self.body())
        with server.connect() as db:before=copy.deepcopy(source_archive.original(db,saved['id']))
        status,_=self.request('/api/transactions/'+saved['id'],{'description':'CORRECTED MANUAL SHOP','amountCents':1825,'status':'approved'})
        self.assertEqual(status,200)
        with server.connect() as db:
            self.assertEqual(source_archive.original(db,saved['id']),before)
            self.assertTrue(db.execute('SELECT 1 FROM transaction_history WHERE transaction_id=? AND action=?',(saved['id'],'manual_entry')).fetchone())
        status,info=self.handler.explain_transaction(saved['id'])
        self.assertEqual(status,200)
        self.assertEqual((info['original']['normalized']['amountCents'],info['current']['amountCents']),(1525,1825))

    def test_later_statement_duplicate_cannot_automatically_double_count(self):
        body=self.body(classification='business',businessId=self.fixture.business,category='Office supplies',includeInTotals=True,status='approved')
        _,saved=self.request('/api/transactions',body)
        with server.connect() as db:
            automation.save(db,{'type':'expense','patterns':['MANUAL SHOP'],'accountId':'card',
                'businessId':self.fixture.business,'category':'Meals','mode':'review','includeInTotals':True},server.utc_now())
        self.fixture.import_csv('card',b'Date,Description,Debit,Credit\n10/22/2026,MANUAL SHOP,15.25,\n')
        with server.connect() as db:
            rows=db.execute("SELECT * FROM transactions WHERE description='MANUAL SHOP'").fetchall()
            self.assertEqual(len(rows),2)
            self.assertEqual(len({r['source_fingerprint'] for r in rows}),1)
            self.assertEqual(sum(r['include_in_totals'] for r in rows),1)
        status,result=self.request('/api/transactions',self.body(**{k:v for k,v in body.items() if k!='requestId'}))
        self.assertEqual(status,400)
        self.assertIn('matching transaction',result['error'])

    def test_manual_rows_are_in_summary_transaction_exports_with_provenance(self):
        _,saved=self.request('/api/transactions',self.body())
        row=self.row(saved['id'])|{'nickname':'Example card','last_four':'','filename':None,'business_name':''}
        values=dict(zip(reports.TRANSACTION_HEADERS,reports.transaction_values(row)))
        self.assertEqual(values['Source document'],'Manual entry')
        self.assertEqual(values['Transaction ID'],saved['id'])
        self.assertEqual(values['Amount USD'],15.25)
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(reports.excel_report([row],'Example business','2026')),data_only=True)
        try:
            self.assertEqual(workbook['Review Queue'].cell(4,13).value,'Manual entry')
        finally:workbook.close()

    def test_included_manual_entry_retains_its_label_in_legacy_csv_ledger(self):
        self.request('/api/transactions',self.body(classification='business',businessId=self.fixture.business,
            category='Office supplies',includeInTotals=True,status='approved'))
        self.handler.path='/api/export.csv';self.handler.wfile=io.BytesIO()
        self.handler.send_response=lambda *args:None;self.handler.send_header=lambda *args:None
        self.handler.end_headers=lambda:None
        self.handler.export_csv()
        rows=list(csv.DictReader(io.StringIO(self.handler.wfile.getvalue().decode('utf-8-sig'))))
        manual=next(row for row in rows if row['Description']=='MANUAL SHOP')
        self.assertEqual(manual['Source document'],'Manual entry')
        self.assertEqual(manual['Amount USD'],'15.25')

