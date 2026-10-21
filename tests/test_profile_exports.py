"""Portable exports reopen complete books and keep import automation working."""
import base64
import io
import sqlite3
import unittest
from pathlib import Path
from unittest.mock import patch
from contextlib import closing

from openpyxl import load_workbook
import automation
import backup
import reports
import server
import statement_import
import test_credits
import workspace_store


class ProfileExportTests(unittest.TestCase):
    def setUp(self):
        self.addCleanup(setattr,server,'PROFILE_PATH',server.PROFILE_PATH)
        self.fixture=test_credits.CreditReviewTests()
        self.fixture.setUp(); self.addCleanup(self.fixture.tearDown)
        self.root=self.fixture.root
        self.handler=self.fixture.handler
        self.handler.download=lambda data,name,kind:(data,name,kind)
        self.config_patch=patch.object(workspace_store,'CONFIG_PATH',self.root/'settings.json')
        self.config_patch.start(); self.addCleanup(self.config_patch.stop)
        self.app_patch=patch.object(server,'ROOT',self.root/'application')
        self.app_patch.start(); self.addCleanup(self.app_patch.stop)
        with server.connect() as db:
            automation.save(db,{'type':'expense','patterns':['SUPPLIER PURCHASE','PLUS MEMBERSHIP'],
                'accountId':'card','businessId':self.fixture.business,'category':'Meals','mode':'review','includeInTotals':True},server.utc_now())
            workspace_store.save_profile(db,server.DB_PATH,name='Portable example')
            workspace_store.save_profile_settings(db,{'sort':'merchant','groupImportsBySourceCategory':True,'workspace':'all','year':'all'})
            db.execute("INSERT INTO import_profiles VALUES('layout','card','example-layout','Example','csv','{}',?,1)",(server.utc_now(),))
            db.execute("INSERT INTO businesses(id,legal_name,display_name,created_at) VALUES('other','Other books','Other books',?)",(server.utc_now(),))
            db.execute("INSERT INTO accounts(id,nickname,business_id) VALUES('other-account','Another account','other')")
        self.fixture.apply({'ids':[self.fixture.ids['Platinum Walmart+ Credit']],'treatment':'statement_credit',
            'rememberRule':True,'rulePatterns':['Walmart'],'ruleAllAccounts':True,'ruleMode':'review'})
        self.fixture.import_csv('other-account',b'Date,Description,Debit,Credit\n10/01/2025,OTHER BUSINESS,25.00,\n')

    def export_excel(self, portable=True):
        self.handler.path='/api/export.xlsx?businessId='+self.fixture.business+'&year=2026&portable='+('1' if portable else '0')
        content,_,_=self.handler.export_excel()
        return content

    def table_state(self):
        with server.connect() as db:
            return {table:[tuple(r) for r in db.execute('SELECT * FROM '+table+' ORDER BY rowid')] for table in (
                'transactions','merchant_rules','credit_rules','accounts','businesses','categories','import_profiles',
                'original_transactions','source_rows','source_archives','transaction_history','review_actions')}

    def reopen(self, content, suffix, upload=False):
        source=self.root/('export'+suffix); source.write_bytes(content)
        target=self.root/'resumed-books'
        body={'mode':'restore','directory':str(target),'sourceDirectory':str(source)}
        if upload: body['contentBase64']=base64.b64encode(content).decode()
        # Simulate a fresh process/session with no books open.
        server.WORKSPACE_READY=False; server.WORKSPACE_ID=None
        self.handler.local_request=lambda:True
        self.handler.path='/api/setup'
        self.handler.read_json=lambda:body
        status,_=self.handler.post_request()
        self.assertEqual(status,200)
        return target

    def assert_future_automation(self):
        self.fixture.import_csv('card',b'Date,Description,Debit,Credit\n11/01/2026,SUPPLIER PURCHASE,17.25,\n11/02/2026,Platinum Walmart+ Credit,,9.50\n')
        with server.connect() as db:
            charge=db.execute("SELECT * FROM transactions WHERE date='2026-11-01'").fetchone()
            credit=db.execute("SELECT * FROM transactions WHERE date='2026-11-02'").fetchone()
            self.assertEqual((charge['category_name'],charge['status'],charge['include_in_totals']),('Meals','approved',1))
            self.assertEqual((credit['credit_treatment'],credit['status'],credit['include_in_totals']),('statement_credit','approved',0))

    def test_sqlite_export_resumes_every_table_rules_and_future_imports(self):
        before=self.table_state()
        payload=backup.export_database(server.DB_PATH)
        target=self.reopen(payload,'.sqlite3',upload=True)
        self.assertEqual(self.table_state(),before)
        self.assertTrue((target/'bookkeeper.profile.json').is_file())
        self.assertEqual(server.active_profile()['name'],'Portable example')
        self.assertEqual(server.active_profile()['settings']['sort'],'merchant')
        self.assert_future_automation()

    def test_filtered_excel_contains_full_profile_and_can_resume_new_session(self):
        before=self.table_state()
        content=self.export_excel()
        workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        try:
            self.assertEqual(workbook[reports.PROFILE_SHEET].sheet_state,'veryHidden')
            self.assertIn('Profile Restore',workbook.sheetnames)
            rows=list(workbook['Expense Ledger'].iter_rows(min_row=4,values_only=True))
            self.assertNotIn('OTHER BUSINESS',{row[1] for row in rows})
        finally: workbook.close()
        self.reopen(content,'.xlsx',upload=True)
        self.assertEqual(self.table_state(),before)
        self.assertEqual(server.active_profile()['settings']['groupImportsBySourceCategory'],True)
        self.assertEqual(server.active_profile()['ruleCount'],2)
        self.assert_future_automation()

    def test_restored_excel_profile_survives_another_session_open(self):
        target=self.reopen(self.export_excel(),'.xlsx')
        saved,_=workspace_store.load_preferences()
        location=workspace_store.locate_books(saved['profile'])
        server.activate_workspace(location['directory'],database=location['database'],profile=location['profile'])
        self.assertEqual(server.DB_PATH,target/'bookkeeper.sqlite3')
        self.assertEqual(server.active_profile()['ruleCount'],2)
        self.assert_future_automation()

    def test_plain_ledger_can_be_shared_without_any_profile_payload(self):
        content=self.export_excel(portable=False)
        workbook=load_workbook(io.BytesIO(content),read_only=True)
        try:
            self.assertNotIn(reports.PROFILE_SHEET,workbook.sheetnames)
            self.assertEqual(len(workbook.sheetnames),7)
        finally:workbook.close()
        with self.assertRaisesRegex(ValueError,'does not contain a reusable profile'):
            backup.restore_backup(io.BytesIO(content),self.root/'plain-restore')
        self.assertFalse((self.root/'plain-restore').exists())

    def test_corrupted_embedded_profile_does_not_create_or_change_books(self):
        workbook=load_workbook(io.BytesIO(self.export_excel()))
        workbook[reports.PROFILE_SHEET]['B3']='AAAA'
        changed=io.BytesIO(); workbook.save(changed); workbook.close()
        before=self.table_state()
        with self.assertRaisesRegex(ValueError,'checksum'):
            backup.restore_backup(io.BytesIO(changed.getvalue()),self.root/'corrupt-restore')
        self.assertFalse((self.root/'corrupt-restore').exists())
        self.assertEqual(self.table_state(),before)

    def test_editing_visible_report_cells_does_not_change_the_profile_snapshot(self):
        before=self.table_state()
        workbook=load_workbook(io.BytesIO(self.export_excel()))
        workbook['Expense Ledger']['E4']='Edited in Excel'
        changed=io.BytesIO();workbook.save(changed);workbook.close()
        self.reopen(changed.getvalue(),'.xlsx')
        self.assertEqual(self.table_state(),before)

    def test_portable_ledger_cannot_be_mistaken_for_a_bank_statement(self):
        with self.assertRaisesRegex(ValueError,'Restore export / backup'):
            statement_import.parse_statement(self.export_excel(),'saved-ledger.xlsx',server.parse_csv)

    def test_restore_cannot_replace_an_existing_database_or_other_sqlite_file(self):
        before=self.table_state()
        with self.assertRaisesRegex(ValueError,'never overwritten'):
            backup.restore_backup(io.BytesIO(self.export_excel()),self.root,allow_empty=True)
        unrelated=self.root/'unrelated.sqlite3'
        with closing(sqlite3.connect(unrelated)) as db:db.execute('CREATE TABLE unrelated (id)')
        with self.assertRaisesRegex(ValueError,'not a Local Bookkeeper database'):
            backup.restore_backup(unrelated,self.root/'unrelated-restore')
        self.assertFalse((self.root/'unrelated-restore').exists())
        self.assertEqual(self.table_state(),before)


if __name__=='__main__':unittest.main()
